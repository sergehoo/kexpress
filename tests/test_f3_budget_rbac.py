"""F3 — Budgets : RBAC, cloisonnement des filiales, approbation, alertes, export, paramètres.

Les garanties F3 tenues ici, chacune par un test qui échouerait si on la retirait :

1. aucun montant budgétaire pour un demandeur ou un chauffeur : 403 sur CHAQUE point budget
   (liste, détail, lignes, révisions, tableau de bord, export, écritures) et aucun montant
   témoin dans aucune réponse — même avec toutes les permissions cochées dans un groupe ;
2. aucune fuite entre filiales : la Finance, l'administrateur ou le gestionnaire de flotte
   d'une filiale sœur ne voit pas les budgets de A (liste vide, détail 404), n'y écrit rien
   et le tableau de bord filtré sur A lui répond 403 ;
3. un budget de GROUPE (filiale vide) n'est lu que par la lecture groupe (administrateur
   entreprise, auditeur, Finance groupe) et n'est créé / modifié que par les écrivains groupe ;
4. l'auditeur lit tout et n'écrit rien (403, base intacte), même par un groupe Django ;
5. approbation : jamais par l'auteur (400 motivé), jamais par la Finance de filiale (403 :
   geste de niveau groupe) ; l'administrateur entreprise et la Finance groupe approuvent ;
   l'archivage suit la même règle ;
6. contrôle positif : la Finance de filiale gère le budget de SA filiale (création, lignes,
   révisions historisées) ;
7. alertes : un seuil franchi par une ligne d'un budget APPROUVÉ crée une alerte UNE fois et
   notifie les financiers et gestionnaires de la filiale seulement ; un brouillon n'alerte
   jamais ; la validation d'une dépense déclenche le contrôle après commit ; les seuils de
   la ligne priment sur ceux du budget, qui priment sur les paramètres Finance ;
8. l'export CSV / XLSX porte les lignes et leurs montants pour un profil habilité, 403 sans
   `export_financial_reports` ;
9. les seuils d'alerte par défaut se paramètrent au niveau GROUPE seulement.

Les budgets sont ceux de l'exercice EN COURS : le contrôle des alertes ne porte que sur lui,
et les dépenses témoins sont datées du jour (mois ouvert, jamais tardives).
"""
import csv
import io
from decimal import Decimal
from types import SimpleNamespace

import pytest
from django.contrib.auth.backends import ModelBackend
from django.contrib.auth.models import Group, Permission
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.core.enums import NotificationType, RoleChoices
from apps.expenses.models import Expense
from apps.expenses.workflow import record_creation
from apps.finance import permissions as perms
from apps.finance.budget import add_line, approve, check_alerts, create_budget
from apps.finance.models import (
    Budget, BudgetAlert, BudgetLine, BudgetRevision, CostCenter, FinanceSettings,
)
from apps.notifications.models import Notification
from tests.test_finance_f1 import _user

pytestmark = pytest.mark.django_db

#: Montants témoins (7 chiffres) : leur présence dans une réponse = un montant servi.
SENTINEL = Decimal("7654321.00")
SENTINEL_TEXT = b"7654321"
G_SENTINEL = Decimal("9182736.00")
G_SENTINEL_TEXT = b"9182736"
ALL_CODENAMES = frozenset(codename for codename, _ in perms.PERMISSIONS)

BUDGETS = "/api/finance/budgets/"
LINES = "/api/finance/budget-lines/"
DASHBOARD = "/api/finance/budgets/dashboard/"
EXPORT = "/api/finance/budgets/export/"
SETTINGS = "/api/finance/settings/"


# --- Outillage -----------------------------------------------------------------------


@pytest.fixture
def api():
    return APIClient()


def _year():
    return timezone.localdate().year


def _rows(response):
    body = response.json()
    return body["results"] if isinstance(body, dict) and "results" in body else body


def _ids(response):
    return {str(row["id"]) for row in _rows(response)}


def _send(api, method, url, data=None):
    if method == "get":
        return api.get(url)
    if method == "delete":
        return api.delete(url)
    return getattr(api, method)(url, data or {}, format="json")


def _grant_by_group(user, codenames):
    """Accorde des permissions `finance.*` par un groupe Django (voie d'exception) et
    renvoie l'utilisateur rechargé (caches de permissions vides)."""
    group = Group.objects.create(name=f"budget-exception-{user.pk}")
    group.permissions.add(*Permission.objects.filter(content_type__app_label="finance",
                                                     codename__in=codenames))
    user.groups.add(group)
    return User.objects.get(pk=user.pk)


def _state():
    """Photographie de tout ce qu'une écriture budgétaire pourrait toucher."""
    return (
        list(Budget.objects.order_by("pk").values_list(
            "pk", "name", "status", "subsidiary_id", "alert_thresholds", "approved_by_id", "year")),
        list(BudgetLine.objects.order_by("pk").values_list(
            "pk", "budget_id", "amount", "month", "subsidiary_id", "cost_center_id", "category", "label",
            "alert_thresholds")),
        list(BudgetRevision.objects.order_by("pk").values_list("pk", "line_id", "new_amount", "kind")),
        BudgetAlert.objects.count(),
        FinanceSettings.current().budget_alert_thresholds,
    )


def _expense(sub, author, amount, *, category="toll", status="validated", label="Péage"):
    """Dépense du jour (mois ouvert) créée par l'ORM avec son auteur et sa ligne de création."""
    expense = Expense.objects.create(
        subsidiary=sub, category=category, label=label, amount=Decimal(amount),
        date=timezone.localdate(), status=status, created_by=author)
    record_creation(expense, author)
    return expense


def _approved_budget(sub, author, approver, lines, *, thresholds=None, name="Budget suivi"):
    budget = create_budget(actor=author, year=_year(), name=name, subsidiary=sub, alert_thresholds=thresholds)
    created = [add_line(budget, actor=author, **spec) for spec in lines]
    approve(budget, actor=approver)
    budget.refresh_from_db()
    return budget, created


def _alert_recipients():
    return set(Notification.objects.filter(notification_type=NotificationType.BUDGET_ALERT)
               .values_list("recipient_id", flat=True))


# --- Profils ----------------------------------------------------------------------------


@pytest.fixture
def fin_a(sub_a):
    return _user("fin-a-f3r@test.io", RoleChoices.FINANCE, sub_a)


@pytest.fixture
def fin_a2(sub_a):
    """Seconde Finance de A : jamais auteur — son refus d'approuver tient au niveau groupe,
    pas à la séparation auteur / approbateur."""
    return _user("fin-a2-f3r@test.io", RoleChoices.FINANCE, sub_a)


@pytest.fixture
def fin_b(sub_b):
    return _user("fin-b-f3r@test.io", RoleChoices.FINANCE, sub_b)


@pytest.fixture
def admin_b(sub_b):
    return _user("admin-b-f3r@test.io", RoleChoices.SUBSIDIARY_ADMIN, sub_b)


@pytest.fixture
def fleet_b(sub_b):
    return _user("fleet-b-f3r@test.io", RoleChoices.FLEET_MANAGER, sub_b)


@pytest.fixture
def group_finance(db):
    return _user("fin-g-f3r@test.io", RoleChoices.FINANCE)


@pytest.fixture
def auditor(db):
    return _user("audit-f3r@test.io", RoleChoices.AUDITOR)


@pytest.fixture
def driver_user(sub_a):
    return _user("drv-f3r@test.io", RoleChoices.DRIVER, sub_a)


@pytest.fixture
def cost_center_a(sub_a):
    return CostCenter.objects.create(subsidiary=sub_a, code="CC-A-F3", name="Siège Abidjan")


@pytest.fixture
def world(sub_a, sub_b, fin_a, fin_b, company_admin):
    """Trois budgets brouillons de l'exercice : A (ligne témoin), B, et GROUPE — dont la ligne
    vise la filiale A : si la lecture groupe fuyait vers A, elle apparaîtrait chez A."""
    year, month = _year(), timezone.localdate().month
    budget_a = create_budget(actor=fin_a, year=year, name="Budget Abidjan", subsidiary=sub_a)
    line_a = add_line(budget_a, actor=fin_a, amount=SENTINEL, month=month, category="toll",
                      label="Péages Abidjan")
    budget_b = create_budget(actor=fin_b, year=year, name="Budget Dakar", subsidiary=sub_b)
    line_b = add_line(budget_b, actor=fin_b, amount=Decimal("1000"), category="toll", label="Péages Dakar")
    budget_g = create_budget(actor=company_admin, year=year, name="Budget Groupe", subsidiary=None)
    line_g = add_line(budget_g, actor=company_admin, amount=G_SENTINEL, subsidiary_id=sub_a.pk,
                      category="energy", label="Énergie groupe")
    return SimpleNamespace(year=year, month=month, a=budget_a, la=line_a, b=budget_b, lb=line_b,
                           g=budget_g, lg=line_g, sub_a=sub_a, sub_b=sub_b)


def _reads(w):
    """(nom, URL, porte le montant témoin de A pour un profil habilité de A ?)."""
    return [
        ("list", BUDGETS, True),
        ("detail", f"{BUDGETS}{w.a.pk}/", True),
        ("lines", LINES, True),
        ("line", f"{LINES}{w.la.pk}/", True),
        ("revisions", f"{LINES}{w.la.pk}/revisions/", True),
        ("dashboard", f"{DASHBOARD}?year={w.year}", True),
        ("dashboard_a", f"{DASHBOARD}?year={w.year}&subsidiary={w.sub_a.pk}", True),
        ("export_csv", f"{EXPORT}?year={w.year}&fmt=csv", True),
        ("export_xlsx", f"{EXPORT}?year={w.year}&fmt=xlsx", False),  # binaire compressé
    ]


def _writes(w):
    """Chaque point d'écriture budgétaire, dans un ordre qu'un écrivain groupe peut dérouler."""
    return [
        ("create", "post", BUDGETS, {"year": w.year, "name": "Intrus", "subsidiary": str(w.sub_a.pk)}),
        ("create_group", "post", BUDGETS, {"year": w.year, "name": "Intrus groupe"}),
        ("rename", "patch", f"{BUDGETS}{w.a.pk}/", {"name": "Renommé"}),
        ("thresholds", "patch", f"{BUDGETS}{w.a.pk}/", {"alert_thresholds": [10]}),
        ("add_line", "post", f"{BUDGETS}{w.a.pk}/lines/", {"amount": "10", "category": "parking"}),
        ("revise", "patch", f"{LINES}{w.la.pk}/", {"amount": "1", "reason": "Intrusion"}),
        ("remove", "delete", f"{LINES}{w.la.pk}/", None),
        ("approve", "post", f"{BUDGETS}{w.a.pk}/approve/", {}),
        ("archive", "post", f"{BUDGETS}{w.a.pk}/archive/", {}),
        ("approve_group", "post", f"{BUDGETS}{w.g.pk}/approve/", {}),
        ("rename_group", "patch", f"{BUDGETS}{w.g.pk}/", {"name": "Renommé"}),
        ("add_line_group", "post", f"{BUDGETS}{w.g.pk}/lines/",
         {"amount": "10", "category": "parking", "reason": "Ajout motivé"}),
    ]


# =====================================================================================
# 1. Aucun montant pour un demandeur ou un chauffeur
# =====================================================================================


def test_1_requester_and_driver_get_403_and_no_amount_on_every_budget_endpoint(
        api, world, fin_a, requester_a, driver_user):
    """Règle absolue : demandeur et chauffeur reçoivent 403 sur chaque point budget, lecture
    comme écriture, sans jamais le montant témoin — même quand un groupe Django leur coche
    toutes les permissions financières — et rien n'est écrit."""
    # Contrôle positif : le témoin est servi à la Finance de A ; son absence plus bas prouve
    # donc un refus, pas une donnée manquante.
    api.force_authenticate(fin_a)
    for name, url, carries in _reads(world):
        response = api.get(url)
        assert response.status_code == 200, f"finance {name} → {response.status_code}"
        if carries:
            assert SENTINEL_TEXT in response.content, f"témoin absent pour la Finance : {name}"

    requester_by_group = _grant_by_group(_user("req-grp-f3r@test.io", RoleChoices.REQUESTER, world.sub_a),
                                         ALL_CODENAMES)
    assert ModelBackend().has_perm(requester_by_group, f"finance.{perms.VIEW_BUDGETS}"), "groupe sans objet"
    before = _state()
    for user in (requester_a, driver_user, requester_by_group):
        api.force_authenticate(user)
        for name, url, _ in _reads(world):
            response = api.get(url)
            assert response.status_code == 403, f"{user.email} {name} → {response.status_code}"
            assert SENTINEL_TEXT not in response.content, f"montant servi à {user.email} : {name}"
            assert G_SENTINEL_TEXT not in response.content
        for name, method, url, data in _writes(world):
            response = _send(api, method, url, data)
            assert response.status_code == 403, f"{user.email} {name} → {response.status_code}"
            assert SENTINEL_TEXT not in response.content and G_SENTINEL_TEXT not in response.content
        assert api.put(SETTINGS, {"budget_alert_thresholds": [10]}, format="json").status_code == 403
    assert _state() == before


# =====================================================================================
# 2. Aucune fuite entre filiales
# =====================================================================================


@pytest.mark.parametrize("profile", ["fin_b", "admin_b", "fleet_b"])
def test_2_sister_subsidiary_neither_sees_nor_writes_budgets_of_a(
        request, api, world, admin_a, cost_center_a, profile):
    """Un profil de la filiale B ne voit aucun budget de A (liste vide de A, détail 404,
    tableau de bord filtré sur A 403, aucun montant témoin), n'y ajoute ni ne révise ni
    n'approuve de ligne, et ne glisse pas une ligne de A dans son propre budget."""
    user = request.getfixturevalue(profile)
    manages = profile != "fleet_b"  # le gestionnaire de flotte lit les budgets, sans les gérer
    hidden_write = 404 if manages else 403
    w = world

    # Contrôle positif : les mêmes URL répondent à un profil de A.
    api.force_authenticate(admin_a)
    assert api.get(f"{BUDGETS}{w.a.pk}/").status_code == 200
    assert api.get(f"{LINES}{w.la.pk}/revisions/").status_code == 200

    before = _state()
    api.force_authenticate(user)
    assert _ids(api.get(BUDGETS)) == {str(w.b.pk)}
    assert _ids(api.get(f"{BUDGETS}?subsidiary={w.sub_a.pk}")) == set()
    assert _ids(api.get(LINES)) == {str(w.lb.pk)}
    for url in (f"{BUDGETS}{w.a.pk}/", f"{LINES}{w.la.pk}/", f"{LINES}{w.la.pk}/revisions/"):
        response = api.get(url)
        assert response.status_code == 404, f"{profile} {url} → {response.status_code}"
    for url in (f"{DASHBOARD}?year={w.year}&subsidiary={w.sub_a.pk}",
                f"{EXPORT}?year={w.year}&fmt=csv&subsidiary={w.sub_a.pk}"):
        response = api.get(url)
        assert response.status_code == 403, f"{profile} {url} → {response.status_code}"
        assert SENTINEL_TEXT not in response.content
    # Sans filtre, ou en citant le budget de A : seulement B, jamais le témoin.
    dashboard = api.get(f"{DASHBOARD}?year={w.year}")
    assert dashboard.status_code == 200
    assert {row["id"] for row in dashboard.json()["lines"]} == {w.lb.pk}
    assert SENTINEL_TEXT not in dashboard.content and G_SENTINEL_TEXT not in dashboard.content
    cited = api.get(f"{DASHBOARD}?year={w.year}&budget={w.a.pk}")
    assert cited.status_code == 200 and cited.json()["lines"] == [] and SENTINEL_TEXT not in cited.content
    export = api.get(f"{EXPORT}?year={w.year}&fmt=csv")
    assert export.status_code == 200 and SENTINEL_TEXT not in export.content

    writes = [
        ("create_for_a", "post", BUDGETS, {"year": w.year, "name": "Intrus", "subsidiary": str(w.sub_a.pk)}, 403),
        ("rename", "patch", f"{BUDGETS}{w.a.pk}/", {"name": "Renommé"}, hidden_write),
        ("add_line", "post", f"{BUDGETS}{w.a.pk}/lines/", {"amount": "10", "category": "parking"}, hidden_write),
        ("revise", "patch", f"{LINES}{w.la.pk}/", {"amount": "1", "reason": "Intrusion"}, hidden_write),
        ("remove", "delete", f"{LINES}{w.la.pk}/", None, hidden_write),
        ("approve", "post", f"{BUDGETS}{w.a.pk}/approve/", {}, 403),
        ("archive", "post", f"{BUDGETS}{w.a.pk}/archive/", {}, 403),
        # Son propre budget, mais une ligne visant la filiale A ou un centre de coût de A.
        ("own_line_for_a", "post", f"{BUDGETS}{w.b.pk}/lines/",
         {"amount": "10", "category": "parking", "subsidiary": str(w.sub_a.pk)}, 400 if manages else 403),
        ("own_line_cc_a", "post", f"{BUDGETS}{w.b.pk}/lines/",
         {"amount": "10", "category": "washing", "cost_center": cost_center_a.pk}, 400 if manages else 403),
    ]
    for name, method, url, data, expected in writes:
        response = _send(api, method, url, data)
        assert response.status_code == expected, f"{profile} {name} → {response.status_code}"
        assert SENTINEL_TEXT not in response.content
    assert _state() == before


# =====================================================================================
# 3. Budgets de GROUPE
# =====================================================================================


def test_3a_group_budget_is_read_by_group_scope_only(
        api, world, company_admin, auditor, group_finance, fin_a, admin_a, fleet_a, fin_b):
    """Un budget de groupe (filiale vide) agrège les filiales sœurs : l'administrateur
    entreprise, l'auditeur et la Finance groupe le lisent ; aucun profil de filiale — pas même
    celle que vise sa ligne — ne le voit ni n'en lit le montant."""
    w = world
    for user in (company_admin, auditor, group_finance):
        api.force_authenticate(user)
        assert str(w.g.pk) in _ids(api.get(BUDGETS)), user.email
        for url in (f"{BUDGETS}{w.g.pk}/", f"{LINES}{w.lg.pk}/", f"{LINES}{w.lg.pk}/revisions/",
                    f"{DASHBOARD}?year={w.year}", f"{EXPORT}?year={w.year}&fmt=csv"):
            response = api.get(url)
            assert response.status_code == 200, f"{user.email} {url} → {response.status_code}"
            assert G_SENTINEL_TEXT in response.content, f"{user.email} ne lit pas le groupe : {url}"
    for user in (fin_a, admin_a, fleet_a, fin_b):
        api.force_authenticate(user)
        assert str(w.g.pk) not in _ids(api.get(BUDGETS)), user.email
        assert str(w.lg.pk) not in _ids(api.get(LINES)), user.email
        for url in (f"{BUDGETS}{w.g.pk}/", f"{LINES}{w.lg.pk}/", f"{LINES}{w.lg.pk}/revisions/"):
            response = api.get(url)
            assert response.status_code == 404, f"{user.email} {url} → {response.status_code}"
        for url in (f"{DASHBOARD}?year={w.year}", f"{DASHBOARD}?year={w.year}&budget={w.g.pk}",
                    f"{EXPORT}?year={w.year}&fmt=csv"):
            response = api.get(url)
            assert response.status_code == 200, f"{user.email} {url} → {response.status_code}"
            assert G_SENTINEL_TEXT not in response.content, f"montant groupe servi à {user.email} : {url}"
    # Contrôle positif : la ligne de A, elle, est bien servie à A (même tableau de bord).
    api.force_authenticate(fin_a)
    assert SENTINEL_TEXT in api.get(f"{DASHBOARD}?year={w.year}").content


def test_3b_only_group_writers_create_or_edit_a_group_budget(
        api, world, company_admin, group_finance, fin_a, admin_a, sub_a):
    """Seuls les écrivains groupe (administrateur entreprise, Finance groupe) créent ou
    modifient un budget de groupe. Un profil de filiale qui omet la filiale crée un budget de
    SA filiale — jamais un budget de groupe — et n'atteint pas le budget de groupe (404)."""
    w = world
    assert Budget.objects.filter(subsidiary__isnull=True).count() == 1
    for user in (fin_a, admin_a):
        api.force_authenticate(user)
        for payload in ({"year": w.year, "name": f"Sans filiale {user.pk}"},
                        {"year": w.year, "name": f"Filiale nulle {user.pk}", "subsidiary": None}):
            response = api.post(BUDGETS, payload, format="json")
            assert response.status_code == 201, response.content
            assert response.json()["subsidiary"] == str(sub_a.pk)
    assert Budget.objects.filter(subsidiary__isnull=True).count() == 1

    # Un administrateur de filiale sans filiale (compte mal rattaché) n'écrit pas au groupe.
    orphan = _user("orphan-f3r@test.io", RoleChoices.SUBSIDIARY_ADMIN)
    api.force_authenticate(orphan)
    assert api.post(BUDGETS, {"year": w.year, "name": "Orphelin"}, format="json").status_code == 403
    assert Budget.objects.filter(subsidiary__isnull=True).count() == 1

    group_before = (list(Budget.objects.filter(pk=w.g.pk).values_list("name", "status", "alert_thresholds")),
                    list(BudgetLine.objects.filter(budget=w.g).order_by("pk").values_list("pk", "amount")),
                    BudgetRevision.objects.filter(line__budget=w.g).count())
    for user in (fin_a, admin_a):
        api.force_authenticate(user)
        for method, url, data in (
                ("patch", f"{BUDGETS}{w.g.pk}/", {"name": "Détourné"}),
                ("post", f"{BUDGETS}{w.g.pk}/lines/", {"amount": "10", "category": "parking"}),
                ("patch", f"{LINES}{w.lg.pk}/", {"amount": "1", "reason": "Détourné"}),
                ("delete", f"{LINES}{w.lg.pk}/", None)):
            response = _send(api, method, url, data)
            assert response.status_code == 404, f"{user.email} {method} {url} → {response.status_code}"
    assert (list(Budget.objects.filter(pk=w.g.pk).values_list("name", "status", "alert_thresholds")),
            list(BudgetLine.objects.filter(budget=w.g).order_by("pk").values_list("pk", "amount")),
            BudgetRevision.objects.filter(line__budget=w.g).count()) == group_before

    # Écrivains groupe : création, renommage, ligne, révision.
    api.force_authenticate(group_finance)
    created = api.post(BUDGETS, {"year": w.year, "name": "Groupe 2"}, format="json")
    assert created.status_code == 201 and created.json()["subsidiary"] is None
    assert Budget.objects.filter(subsidiary__isnull=True).count() == 2
    line = api.post(f"{BUDGETS}{w.g.pk}/lines/", {"amount": "500", "category": "parking"}, format="json")
    assert line.status_code == 201 and line.json()["amount"] == "500.00"
    revised = api.patch(f"{LINES}{w.lg.pk}/", {"amount": "1000"}, format="json")
    assert revised.status_code == 200 and revised.json()["amount"] == "1000.00"
    api.force_authenticate(company_admin)
    renamed = api.patch(f"{BUDGETS}{w.g.pk}/", {"name": "Budget Groupe révisé"}, format="json")
    assert renamed.status_code == 200
    w.g.refresh_from_db()
    assert w.g.name == "Budget Groupe révisé" and w.g.subsidiary_id is None


# =====================================================================================
# 4. L'auditeur lit tout, n'écrit rien
# =====================================================================================


def test_4a_auditor_reads_every_budget(api, world, auditor):
    """L'auditeur (lecture groupe) lit les budgets de toutes les filiales et du groupe :
    liste, détail, lignes, révisions, tableau de bord (filtré ou non), export CSV / XLSX."""
    w = world
    api.force_authenticate(auditor)
    assert {str(w.a.pk), str(w.b.pk), str(w.g.pk)} <= _ids(api.get(BUDGETS))
    for budget in (w.a, w.b, w.g):
        assert api.get(f"{BUDGETS}{budget.pk}/").status_code == 200
    revisions = api.get(f"{LINES}{w.la.pk}/revisions/")
    assert revisions.status_code == 200
    assert [(r["kind"], r["previous_amount"], r["new_amount"]) for r in revisions.json()] == \
        [("initial", None, str(SENTINEL))]
    dashboard = api.get(f"{DASHBOARD}?year={w.year}")
    assert dashboard.status_code == 200
    assert {row["id"] for row in dashboard.json()["lines"]} == {w.la.pk, w.lb.pk, w.lg.pk}
    assert SENTINEL_TEXT in dashboard.content and G_SENTINEL_TEXT in dashboard.content
    assert api.get(f"{DASHBOARD}?year={w.year}&subsidiary={w.sub_b.pk}").status_code == 200
    export = api.get(f"{EXPORT}?year={w.year}&fmt=csv")
    assert export.status_code == 200 and export["Content-Type"].startswith("text/csv")
    assert SENTINEL_TEXT in export.content and G_SENTINEL_TEXT in export.content
    assert api.get(f"{EXPORT}?year={w.year}&fmt=xlsx").status_code == 200


def test_4b_auditor_writes_nothing_even_through_a_django_group(api, world, auditor):
    """D7 : chaque écriture budgétaire de l'auditeur est refusée (403) et la base est
    intacte — y compris quand un groupe Django lui coche `manage_budgets`, `approve_budget`
    et `manage_finance_settings`."""
    granted = _grant_by_group(auditor, [perms.MANAGE_BUDGETS, perms.APPROVE_BUDGET,
                                        perms.MANAGE_FINANCE_SETTINGS])
    for codename in (perms.MANAGE_BUDGETS, perms.APPROVE_BUDGET, perms.MANAGE_FINANCE_SETTINGS):
        assert ModelBackend().has_perm(granted, f"finance.{codename}"), f"groupe sans objet : {codename}"
        assert not perms.can(granted, codename), codename
    before = _state()
    for user in (auditor, granted):
        api.force_authenticate(user)
        for name, method, url, data in _writes(world):
            response = _send(api, method, url, data)
            assert response.status_code == 403, f"auditeur {name} → {response.status_code}"
        assert api.put(SETTINGS, {"budget_alert_thresholds": [10]}, format="json").status_code == 403
    assert _state() == before


def test_4c_the_same_writes_succeed_for_a_group_writer(api, world, group_finance):
    """Contrôle positif du 4b : les requêtes refusées à l'auditeur sont bien formées — des
    écrivains groupe les déroulent toutes avec succès. La Finance groupe saisit (elle ajoute
    une ligne à A) : l'approbation et l'archivage de A reviennent donc à un AUTRE profil
    groupe — qui a saisi des montants n'approuve pas ses propres chiffres."""
    other = _user("ca2-f3r@test.io", RoleChoices.COMPANY_ADMIN)
    expected = {"create": 201, "create_group": 201, "rename": 200, "thresholds": 200, "add_line": 201,
                "revise": 200, "remove": 204, "approve": 200, "archive": 200, "approve_group": 200,
                "rename_group": 200, "add_line_group": 201}
    for name, method, url, data in _writes(world):
        api.force_authenticate(other if name in ("approve", "archive") else group_finance)
        response = _send(api, method, url, data)
        assert response.status_code == expected[name], f"{name} → {response.status_code} {response.content}"
    world.a.refresh_from_db()
    world.g.refresh_from_db()
    assert (world.a.status, world.a.approved_by_id) == (Budget.ARCHIVED, other.pk)
    assert (world.g.status, world.g.approved_by_id) == (Budget.APPROVED, group_finance.pk)


# =====================================================================================
# 5. Approbation et archivage : niveau groupe, jamais l'auteur
# =====================================================================================


def test_5a_the_author_never_approves_his_own_budget(api, world, company_admin, group_finance, sub_a):
    """Séparation des responsabilités : l'auteur d'un budget ne l'approuve jamais (400
    motivé, budget inchangé), même administrateur entreprise ou Finance groupe ; une autre
    personne habilitée l'approuve."""
    w = world
    api.force_authenticate(company_admin)  # auteur du budget de groupe
    response = api.post(f"{BUDGETS}{w.g.pk}/approve/")
    assert response.status_code == 400
    assert "autre personne que son auteur" in str(response.json())
    w.g.refresh_from_db()
    assert (w.g.status, w.g.approved_by_id, w.g.approved_at) == (Budget.DRAFT, None, None)

    api.force_authenticate(group_finance)
    own = api.post(BUDGETS, {"year": w.year, "name": "Budget Finance groupe", "subsidiary": str(sub_a.pk)},
                   format="json")
    assert own.status_code == 201
    own_id = own.json()["id"]
    assert api.post(f"{BUDGETS}{own_id}/lines/", {"amount": "100", "category": "toll"},
                    format="json").status_code == 201
    response = api.post(f"{BUDGETS}{own_id}/approve/")
    assert response.status_code == 400 and "autre personne que son auteur" in str(response.json())
    assert Budget.objects.get(pk=own_id).status == Budget.DRAFT

    # Chacun approuve le budget de l'autre.
    approved = api.post(f"{BUDGETS}{w.g.pk}/approve/")
    assert approved.status_code == 200 and approved.json()["status"] == Budget.APPROVED
    api.force_authenticate(company_admin)
    assert api.post(f"{BUDGETS}{own_id}/approve/").status_code == 200
    w.g.refresh_from_db()
    assert (w.g.status, w.g.approved_by_id) == (Budget.APPROVED, group_finance.pk)
    assert Budget.objects.get(pk=own_id).approved_by_id == company_admin.pk


def test_5b_subsidiary_profiles_cannot_approve_or_archive_group_level(
        api, world, fin_a, fin_a2, admin_a, company_admin, group_finance):
    """Approuver et archiver engagent le groupe : la Finance de filiale (même non auteur) et
    l'administrateur de filiale reçoivent 403 sur le budget de LEUR filiale, sans effet ;
    l'administrateur entreprise approuve, la Finance groupe approuve et archive."""
    w = world
    for user in (fin_a2, fin_a, admin_a):
        api.force_authenticate(user)
        response = api.post(f"{BUDGETS}{w.a.pk}/approve/")
        assert response.status_code == 403, f"{user.email} → {response.status_code}"
    w.a.refresh_from_db()
    assert (w.a.status, w.a.approved_by_id) == (Budget.DRAFT, None)

    api.force_authenticate(company_admin)
    assert api.post(f"{BUDGETS}{w.a.pk}/approve/").status_code == 200
    api.force_authenticate(group_finance)
    assert api.post(f"{BUDGETS}{w.b.pk}/approve/").status_code == 200

    for user in (fin_a2, fin_a, admin_a):
        api.force_authenticate(user)
        response = api.post(f"{BUDGETS}{w.a.pk}/archive/")
        assert response.status_code == 403, f"{user.email} archive → {response.status_code}"
    w.a.refresh_from_db()
    assert (w.a.status, w.a.approved_by_id) == (Budget.APPROVED, company_admin.pk)

    api.force_authenticate(group_finance)
    archived = api.post(f"{BUDGETS}{w.a.pk}/archive/")
    assert archived.status_code == 200 and archived.json()["status"] == Budget.ARCHIVED
    api.force_authenticate(company_admin)
    assert api.post(f"{BUDGETS}{w.b.pk}/archive/").status_code == 200
    assert set(Budget.objects.filter(pk__in=(w.a.pk, w.b.pk)).values_list("status", flat=True)) == {Budget.ARCHIVED}


# =====================================================================================
# 6. Contrôle positif : la Finance de filiale gère le budget de sa filiale
# =====================================================================================


def test_6_subsidiary_finance_manages_its_own_subsidiary_budget(
        api, sub_a, fin_a, admin_a, fleet_a, company_admin):
    """La Finance de A crée le budget de A, y ajoute et révise des lignes (librement en
    brouillon, avec motif une fois approuvé) ; l'administrateur de A ajoute une ligne ; le
    gestionnaire de flotte, lecteur, est refusé. Chaque montant est historisé."""
    year = _year()
    api.force_authenticate(fin_a)
    created = api.post(BUDGETS, {"year": year, "name": "Budget A exploitation", "subsidiary": str(sub_a.pk)},
                       format="json")
    assert created.status_code == 201, created.content
    body = created.json()
    assert (body["subsidiary"], body["status"], body["created_by"]) == (str(sub_a.pk), Budget.DRAFT, str(fin_a.pk))
    budget_id = body["id"]
    line = api.post(f"{BUDGETS}{budget_id}/lines/", {"amount": "120000", "category": "parking",
                                                     "label": "Parkings"}, format="json")
    assert line.status_code == 201 and line.json()["amount"] == "120000.00"
    line_id = line.json()["id"]

    api.force_authenticate(admin_a)
    assert api.post(f"{BUDGETS}{budget_id}/lines/", {"amount": "50000", "category": "washing"},
                    format="json").status_code == 201
    api.force_authenticate(fleet_a)
    lines_before = BudgetLine.objects.filter(budget_id=budget_id).count()
    assert api.post(f"{BUDGETS}{budget_id}/lines/", {"amount": "1", "category": "toll"},
                    format="json").status_code == 403
    assert api.patch(f"{LINES}{line_id}/", {"amount": "1"}, format="json").status_code == 403
    assert BudgetLine.objects.filter(budget_id=budget_id).count() == lines_before

    api.force_authenticate(fin_a)
    draft = api.patch(f"{LINES}{line_id}/", {"amount": "150000"}, format="json")
    assert draft.status_code == 200 and draft.json()["amount"] == "150000.00"

    api.force_authenticate(company_admin)
    assert api.post(f"{BUDGETS}{budget_id}/approve/").status_code == 200

    api.force_authenticate(fin_a)
    unmotivated = api.patch(f"{LINES}{line_id}/", {"amount": "160000"}, format="json")
    assert unmotivated.status_code == 400 and "motif" in str(unmotivated.json())
    assert BudgetLine.objects.get(pk=line_id).amount == Decimal("150000.00")
    revised = api.patch(f"{LINES}{line_id}/", {"amount": "160000", "reason": "Hausse des tarifs"}, format="json")
    assert revised.status_code == 200 and revised.json()["amount"] == "160000.00"
    assert api.post(f"{BUDGETS}{budget_id}/lines/", {"amount": "9000", "category": "toll"},
                    format="json").status_code == 400  # approuvé : nouvelle ligne = révision motivée
    added = api.post(f"{BUDGETS}{budget_id}/lines/", {"amount": "9000", "category": "toll",
                                                      "reason": "Nouveau péage"}, format="json")
    assert added.status_code == 201
    assert api.delete(f"{LINES}{line_id}/").status_code == 400  # approuvé : on révise, on ne supprime pas

    history = api.get(f"{LINES}{line_id}/revisions/")
    assert history.status_code == 200
    assert [(r["kind"], r["previous_amount"], r["new_amount"], r["reason"]) for r in history.json()] == [
        ("initial", None, "120000.00", ""),
        ("draft", "120000.00", "150000.00", ""),
        ("revision", "150000.00", "160000.00", "Hausse des tarifs"),
    ]
    assert BudgetRevision.objects.get(line_id=added.json()["id"]).kind == "revision"


# =====================================================================================
# 7. Alertes budgétaires
# =====================================================================================


@pytest.fixture
def audience(fin_a, group_finance, company_admin, admin_a, fleet_a, requester_a, driver_user, manager_a,
             fin_b, fleet_b, admin_b, auditor):
    """Tous les profils présents : les destinataires attendus d'une alerte de A, et ceux qui
    ne doivent jamais la recevoir (demandeur, chauffeur, manager, auditeur, filiale B)."""
    return SimpleNamespace(expected={fin_a.pk, group_finance.pk, company_admin.pk, admin_a.pk, fleet_a.pk},
                           never={requester_a.pk, driver_user.pk, manager_a.pk, fin_b.pk, fleet_b.pk,
                                  admin_b.pk, auditor.pk})


def test_7a_threshold_crossing_alerts_once_and_only_the_subsidiary_audience(
        sub_a, fin_a, fleet_a, company_admin, audience):
    """Seuils [50, 100] : une ligne approuvée consommée à 60 % crée UNE alerte 50 et notifie
    les financiers et gestionnaires de A seulement ; un second contrôle ne duplique rien ;
    franchir 100 % ajoute l'alerte 100, notifiée aux mêmes."""
    budget, (line,) = _approved_budget(sub_a, fin_a, company_admin,
                                       [{"amount": Decimal("10000"), "category": "toll", "label": "Péages"}],
                                       thresholds=[50, 100])
    _expense(sub_a, fleet_a, "6000.00")

    assert check_alerts() == 1
    alert = BudgetAlert.objects.get(line=line)
    assert (alert.threshold, alert.rate, alert.consumed, alert.planned) == \
        (50, Decimal("60.00"), Decimal("6000.00"), Decimal("10000.00"))
    assert _alert_recipients() == audience.expected
    assert not _alert_recipients() & audience.never
    assert Notification.objects.filter(notification_type=NotificationType.BUDGET_ALERT).count() == \
        len(audience.expected)

    assert check_alerts() == 0
    assert BudgetAlert.objects.filter(line=line).count() == 1
    assert Notification.objects.filter(notification_type=NotificationType.BUDGET_ALERT).count() == \
        len(audience.expected)

    _expense(sub_a, fleet_a, "4500.00", label="Péage retour")
    assert check_alerts() == 1
    assert sorted(BudgetAlert.objects.filter(line=line).values_list("threshold", "rate")) == \
        [(50, Decimal("60.00")), (100, Decimal("105.00"))]
    hundred = Notification.objects.filter(notification_type=NotificationType.BUDGET_ALERT, title__contains="100 %")
    assert set(hundred.values_list("recipient_id", flat=True)) == audience.expected
    assert hundred.count() == len(audience.expected)
    assert not _alert_recipients() & audience.never


def test_7b_a_draft_budget_never_alerts(sub_a, fin_a, fleet_a):
    """Un brouillon n'est pas un engagement du groupe : consommé à 200 %, il n'alerte pas —
    même quand on le passe explicitement au contrôle."""
    budget = create_budget(actor=fin_a, year=_year(), name="Brouillon", subsidiary=sub_a, alert_thresholds=[50, 100])
    line = add_line(budget, actor=fin_a, amount=Decimal("1000"), category="toll")
    _expense(sub_a, fleet_a, "2000.00")
    assert check_alerts() == 0
    assert check_alerts(Budget.objects.filter(pk=budget.pk)) == 0
    assert not BudgetAlert.objects.filter(line=line).exists()
    assert not Notification.objects.filter(notification_type=NotificationType.BUDGET_ALERT).exists()


def test_7c_expense_validation_triggers_the_alert_check_after_commit(
        api, sub_a, fin_a, fleet_a, company_admin, django_capture_on_commit_callbacks):
    """La validation d'une dépense programme `check_alerts_for` APRÈS commit : rien n'est
    alerté tant que la transaction n'est pas validée, l'alerte naît à l'exécution du rappel."""
    budget, (line,) = _approved_budget(sub_a, fin_a, company_admin,
                                       [{"amount": Decimal("10000"), "category": "toll"}], thresholds=[50])
    expense = _expense(sub_a, fleet_a, "6000.00", status="to_validate")
    api.force_authenticate(fin_a)
    with django_capture_on_commit_callbacks(execute=False) as callbacks:
        response = api.post(f"/api/expenses/{expense.pk}/validate/", {}, format="json")
    assert response.status_code == 200, response.content
    expense.refresh_from_db()
    assert expense.status == "validated"
    assert callbacks, "aucun rappel après commit programmé par la validation"
    assert not BudgetAlert.objects.exists()

    for callback in callbacks:
        callback()
    alert = BudgetAlert.objects.get(line=line)
    assert (alert.threshold, alert.rate, alert.consumed) == (50, Decimal("60.00"), Decimal("6000.00"))
    assert fin_a.pk in _alert_recipients()


def test_7d_line_thresholds_override_budget_thresholds_which_override_settings(
        sub_a, fin_a, fleet_a, company_admin):
    """Hiérarchie des seuils : ceux de la LIGNE priment, puis ceux du BUDGET, puis les
    paramètres Finance. À 60 % de consommation partout : paramètres [40] → alerte 40 ;
    budget [50] → alerte 50 (pas 40) ; ligne [30, 70] → alerte 30 (ni 40 ni 50)."""
    settings_row = FinanceSettings.current()
    assert settings_row.budget_alert_thresholds == [80, 90, 100]  # défaut : 60 % n'alerterait pas
    settings_row.budget_alert_thresholds = [40]
    settings_row.save()

    _, (by_settings,) = _approved_budget(sub_a, fin_a, company_admin,
                                         [{"amount": Decimal("10000"), "category": "toll"}], name="Paramètres")
    _, (by_budget, by_line) = _approved_budget(
        sub_a, fin_a, company_admin,
        [{"amount": Decimal("10000"), "category": "parking"},
         {"amount": Decimal("10000"), "category": "washing", "alert_thresholds": [30, 70]}],
        thresholds=[50], name="Seuils propres")
    for category in ("toll", "parking", "washing"):
        _expense(sub_a, fleet_a, "6000.00", category=category, label=category)

    assert check_alerts() == 3
    thresholds = {line.pk: sorted(BudgetAlert.objects.filter(line=line).values_list("threshold", flat=True))
                  for line in (by_settings, by_budget, by_line)}
    assert thresholds == {by_settings.pk: [40], by_budget.pk: [50], by_line.pk: [30]}
    assert set(BudgetAlert.objects.values_list("rate", flat=True)) == {Decimal("60.00")}


def test_7e_group_budget_alert_never_reaches_subsidiary_profiles(
        api, sub_a, company_admin, group_finance, fin_a, admin_a, fleet_a):
    """Un budget de groupe est invisible aux profils de filiale (garantie 3) : son alerte —
    même pour une ligne qui vise la filiale A — ne leur est pas notifiée non plus ; elle va à
    la lecture groupe seulement, et ne porte jamais le montant prévu chez A."""
    _, (line,) = _approved_budget(None, company_admin, group_finance,
                                  [{"amount": G_SENTINEL, "subsidiary_id": sub_a.pk, "category": "toll"}],
                                  thresholds=[50], name="Budget Groupe confidentiel")
    _expense(sub_a, fleet_a, "6000000.00")
    api.force_authenticate(fin_a)
    assert api.get(f"{BUDGETS}{line.budget_id}/").status_code == 404  # invisible à A

    assert check_alerts() == 1
    recipients = _alert_recipients()
    assert group_finance.pk in recipients  # contrôle positif : l'alerte existe et part au groupe
    assert not recipients & {fin_a.pk, admin_a.pk, fleet_a.pk}, "budget de groupe notifié à la filiale"
    assert not Notification.objects.filter(recipient__in=(fin_a, admin_a, fleet_a),
                                           message__contains="9182736").exists()


# =====================================================================================
# 8. Export
# =====================================================================================


def test_8a_export_csv_and_xlsx_carry_the_lines_and_their_amounts(api, world, fin_a, fleet_a):
    """Un profil habilité exporte ses lignes avec prévu, engagé, réalisé, décaissé et
    disponible — en CSV comme en XLSX — et seulement celles de son périmètre."""
    from openpyxl import load_workbook

    _expense(world.sub_a, fleet_a, "1234.00")
    api.force_authenticate(fin_a)

    response = api.get(f"{EXPORT}?year={world.year}&fmt=csv")
    assert response.status_code == 200 and response["Content-Type"].startswith("text/csv")
    assert f'filename="budget_{world.year}.csv"' in response["Content-Disposition"]
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    header, data = rows[0], rows[1:]
    assert header[:6] == ["Budget", "Statut", "Mois", "Filiale", "Centre de coût", "Catégorie"]
    assert len(data) == 1, data  # ni B, ni le groupe
    row = dict(zip(header, data[0]))
    assert (row["Budget"], row["Statut"], row["Mois"], row["Filiale"], row["Centre de coût"], row["Catégorie"]) \
        == ("Budget Abidjan", "Brouillon", str(world.month), "Abidjan", "Tous", "Péage")
    assert (row["Prévu"], row["Engagé"], row["Réalisé"], row["Décaissé"], row["Disponible"]) == \
        ("7654321.00", "0.00", "1234.00", "0.00", "7653087.00")

    response = api.get(f"{EXPORT}?year={world.year}&fmt=xlsx")
    assert response.status_code == 200
    assert f'filename="budget_{world.year}.xlsx"' in response["Content-Disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == header
    assert len(rows) == 2
    xrow = dict(zip(header, rows[1]))
    assert xrow["Budget"] == "Budget Abidjan" and xrow["Catégorie"] == "Péage"
    assert {k: Decimal(str(xrow[k])) for k in ("Prévu", "Engagé", "Réalisé", "Décaissé", "Disponible")} == {
        "Prévu": SENTINEL, "Engagé": Decimal("0"), "Réalisé": Decimal("1234"), "Décaissé": Decimal("0"),
        "Disponible": Decimal("7653087")}


def test_8b_export_requires_export_financial_reports(api, world, manager_a):
    """Lire les budgets n'est pas les exporter : un profil qui tient `view_budgets` (par un
    groupe Django) sans `export_financial_reports` lit le tableau de bord mais reçoit 403 à
    l'export, CSV comme XLSX, sans le montant témoin."""
    reader = _grant_by_group(manager_a, [perms.VIEW_BUDGETS])
    assert perms.can(reader, perms.VIEW_BUDGETS) and not perms.can(reader, perms.EXPORT_FINANCIAL_REPORTS)
    api.force_authenticate(reader)
    dashboard = api.get(f"{DASHBOARD}?year={world.year}")
    assert dashboard.status_code == 200 and SENTINEL_TEXT in dashboard.content  # contrôle positif
    for fmt in ("csv", "xlsx"):
        response = api.get(f"{EXPORT}?year={world.year}&fmt={fmt}")
        assert response.status_code == 403, f"{fmt} → {response.status_code}"
        assert SENTINEL_TEXT not in response.content


# =====================================================================================
# 9. Seuils d'alerte par défaut : niveau groupe
# =====================================================================================


def test_9_budget_alert_thresholds_are_set_at_group_level_only(
        api, group_finance, company_admin, fin_a, admin_a, auditor, requester_a):
    """`manage_finance_settings` est un geste de niveau groupe : la Finance groupe (et
    l'administrateur entreprise) règlent les seuils par défaut ; la Finance de filiale,
    l'administrateur de filiale, l'auditeur — même par un groupe Django — et le demandeur
    sont refusés (403) et les seuils restent inchangés."""
    assert FinanceSettings.current().budget_alert_thresholds == [80, 90, 100]
    api.force_authenticate(group_finance)
    response = api.put(SETTINGS, {"budget_alert_thresholds": [100, 50, 75]}, format="json")
    assert response.status_code == 200, response.content
    assert response.json()["budget_alert_thresholds"] == [50, 75, 100]
    assert FinanceSettings.current().budget_alert_thresholds == [50, 75, 100]
    invalid = api.put(SETTINGS, {"budget_alert_thresholds": [0, 50]}, format="json")
    assert invalid.status_code == 400
    assert FinanceSettings.current().budget_alert_thresholds == [50, 75, 100]

    auditor_by_group = _grant_by_group(_user("audit2-f3r@test.io", RoleChoices.AUDITOR),
                                       [perms.MANAGE_FINANCE_SETTINGS])
    assert ModelBackend().has_perm(auditor_by_group, f"finance.{perms.MANAGE_FINANCE_SETTINGS}")
    for user in (fin_a, admin_a, auditor, auditor_by_group, requester_a):
        api.force_authenticate(user)
        for method in ("put", "patch"):
            response = getattr(api, method)(SETTINGS, {"budget_alert_thresholds": [10]}, format="json")
            assert response.status_code == 403, f"{user.email} {method} → {response.status_code}"
    assert FinanceSettings.current().budget_alert_thresholds == [50, 75, 100]

    api.force_authenticate(company_admin)
    assert api.patch(SETTINGS, {"budget_alert_thresholds": [90]}, format="json").status_code == 200
    assert FinanceSettings.current().budget_alert_thresholds == [90]
